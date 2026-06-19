#!/usr/bin/env python3
"""Builds a PowerPoint deck on the ASUS ROG Astral GeForce RTX 5080,
focused on how long the card will stay relevant."""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

# ---- palette ---------------------------------------------------------------
INK      = RGBColor(0x10, 0x14, 0x1C)   # near-black background
PANEL    = RGBColor(0x1A, 0x20, 0x2E)   # card panel
ACCENT   = RGBColor(0x35, 0xD0, 0x7A)   # nvidia-ish green
ACCENT2  = RGBColor(0x4C, 0x9A, 0xFF)   # blue
WARN     = RGBColor(0xF0, 0xA5, 0x3A)   # amber
WHITE    = RGBColor(0xF2, 0xF5, 0xF8)
MUTED    = RGBColor(0x9A, 0xA7, 0xB8)
LINE     = RGBColor(0x2C, 0x35, 0x46)

prs = Presentation()
prs.slide_width  = Inches(13.333)
prs.slide_height = Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]


def bg(slide, color=INK):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color


def rect(slide, x, y, w, h, fill=None, line=None, line_w=1.0):
    from pptx.enum.shapes import MSO_SHAPE
    sp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    sp.adjustments[0] = 0.06
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid(); sp.fill.fore_color.rgb = fill
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line; sp.line.width = Pt(line_w)
    sp.shadow.inherit = False
    return sp


def txt(slide, x, y, w, h, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
        space_after=6):
    """runs: list of paragraphs; each paragraph is list of (text,size,color,bold,italic)."""
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for i, para in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space_after)
        p.space_before = Pt(0)
        for (t, s, c, b, it) in para:
            r = p.add_run(); r.text = t
            r.font.size = Pt(s); r.font.color.rgb = c
            r.font.bold = b; r.font.italic = it
            r.font.name = "Segoe UI"
    return tb


def bullets(slide, x, y, w, h, items, size=15, color=WHITE, gap=10, dot=ACCENT):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame; tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(gap); p.space_before = Pt(0)
        r = p.add_run(); r.text = "▪  "
        r.font.size = Pt(size); r.font.color.rgb = dot; r.font.bold = True
        r.font.name = "Segoe UI"
        r2 = p.add_run(); r2.text = item
        r2.font.size = Pt(size); r2.font.color.rgb = color
        r2.font.name = "Segoe UI"
    return tb


def kicker(slide, text, color=ACCENT):
    txt(slide, Inches(0.7), Inches(0.45), Inches(11), Inches(0.4),
        [[(text.upper(), 13, color, True, False)]])


def title(slide, text, y=Inches(0.8), size=34):
    txt(slide, Inches(0.7), y, Inches(12), Inches(1.0),
        [[(text, size, WHITE, True, False)]])


def accent_bar(slide):
    bar = rect(slide, Inches(0.7), Inches(0.78), Inches(0.9), Inches(0.07), fill=ACCENT)
    return bar


# ===========================================================================
# Slide 1 — Title
# ===========================================================================
s = prs.slides.add_slide(BLANK); bg(s)
rect(s, 0, 0, SW, Inches(0.18), fill=ACCENT)
txt(s, Inches(0.9), Inches(2.0), Inches(11.5), Inches(0.5),
    [[("ASUS ROG ASTRAL", 18, ACCENT, True, False)]])
txt(s, Inches(0.9), Inches(2.5), Inches(11.5), Inches(2.0),
    [[("GeForce RTX 5080", 54, WHITE, True, False)]])
txt(s, Inches(0.9), Inches(3.7), Inches(11.5), Inches(1.0),
    [[("How long will it stay relevant?", 26, MUTED, False, True)]])
txt(s, Inches(0.9), Inches(5.6), Inches(11.5), Inches(0.6),
    [[("A buyer's longevity & future-proofing briefing", 15, MUTED, False, False)]])
rect(s, Inches(0.9), Inches(5.35), Inches(3.2), Inches(0.04), fill=LINE)

# ===========================================================================
# Slide 2 — What this card is
# ===========================================================================
s = prs.slides.add_slide(BLANK); bg(s)
kicker(s, "The product"); accent_bar(s)
title(s, "What the ROG Astral RTX 5080 is")
txt(s, Inches(0.7), Inches(1.7), Inches(12), Inches(1.0),
    [[("A flagship-tier partner card: ASUS's top-of-the-stack ROG Astral cooler "
       "wrapped around NVIDIA's GeForce RTX 5080 — the second-fastest GPU of the "
       "Blackwell (RTX 50-series) generation, launched in 2025.", 16, WHITE, False, False)]])

cards = [
    ("GPU", "NVIDIA Blackwell (GB203)\nRTX 5080 silicon", ACCENT),
    ("Memory", "16 GB GDDR7\n256-bit bus, very high bandwidth", ACCENT2),
    ("Key feature", "DLSS 4 with\nMulti-Frame Generation", WARN),
    ("Cooler", "ROG Astral quad-fan\npremium thermal design", ACCENT),
]
cw = Inches(2.95); gap = Inches(0.18); x0 = Inches(0.7); y0 = Inches(3.1)
for i, (h, b, c) in enumerate(cards):
    x = x0 + i * (cw + gap)
    rect(s, x, y0, cw, Inches(2.4), fill=PANEL, line=LINE, line_w=1)
    rect(s, x, y0, Inches(0.12), Inches(2.4), fill=c)
    txt(s, x + Inches(0.25), y0 + Inches(0.25), cw - Inches(0.4), Inches(0.5),
        [[(h.upper(), 12, c, True, False)]])
    txt(s, x + Inches(0.25), y0 + Inches(0.85), cw - Inches(0.4), Inches(1.4),
        [[(b, 15, WHITE, False, False)]])
txt(s, Inches(0.7), Inches(5.9), Inches(12), Inches(0.6),
    [[("Positioning: high-end 4K gaming and creator card — one tier below the "
       "RTX 5090 flagship.", 13, MUTED, False, True)]])

# ===========================================================================
# Slide 3 — Why these specs matter for longevity
# ===========================================================================
s = prs.slides.add_slide(BLANK); bg(s)
kicker(s, "The longevity drivers"); accent_bar(s)
title(s, "Which specs decide how long it lasts")
rows = [
    ("Raw performance", "Sits near the top of the stack — comfortably drives 4K today "
                        "and has headroom to spare, which is the #1 predictor of a long life.", ACCENT),
    ("16 GB VRAM", "Generous for 4K now, but the most likely future bottleneck. Texture "
                   "and ray-tracing memory demands climb fastest of any metric.", WARN),
    ("DLSS 4 + Frame Gen", "Software upside: NVIDIA's upscaling/frame-gen keeps extending "
                           "playable performance long after raw silicon would tap out.", ACCENT2),
    ("Modern I/O", "PCIe 5.0, DisplayPort 2.1, HDMI 2.1 — ready for next-gen monitors and "
                   "platforms, so the rest of your PC won't make it obsolete early.", ACCENT),
]
y = Inches(1.75)
for h, b, c in rows:
    rect(s, Inches(0.7), y, Inches(11.95), Inches(1.15), fill=PANEL, line=LINE, line_w=1)
    rect(s, Inches(0.7), y, Inches(0.12), Inches(1.15), fill=c)
    txt(s, Inches(1.0), y + Inches(0.15), Inches(3.1), Inches(0.9),
        [[(h, 17, c, True, False)]], anchor=MSO_ANCHOR.MIDDLE)
    txt(s, Inches(4.2), y + Inches(0.12), Inches(8.2), Inches(0.95),
        [[(b, 14, WHITE, False, False)]], anchor=MSO_ANCHOR.MIDDLE)
    y += Inches(1.30)

# ===========================================================================
# Slide 4 — Relevance timeline (the core answer)
# ===========================================================================
s = prs.slides.add_slide(BLANK); bg(s)
kicker(s, "The answer", WARN); accent_bar(s)
title(s, "Projected relevance timeline")

# baseline assumes a 2025 purchase
phases = [
    ("2025–2027", "PEAK", "Top-tier 4K, max settings, high frame rates with DLSS. "
        "No compromises.", ACCENT),
    ("2028–2029", "STRONG", "Still excellent at 4K; occasional setting trims on the most "
        "demanding titles. DLSS 4 carries the load.", ACCENT2),
    ("2030–2031", "CAPABLE", "A solid high-end 1440p / entry-4K card. VRAM begins to limit "
        "max textures in newest games.", WARN),
    ("2032+", "AGING", "Very playable at 1080p/1440p with upscaling, but no longer a "
        "high-end card. Upgrade territory.", MUTED),
]
n = len(phases); cw = Inches(2.92); gap = Inches(0.18); x0 = Inches(0.7); y0 = Inches(2.4)
# connecting line
rect(s, x0, y0 - Inches(0.35), Inches(11.95), Inches(0.04), fill=LINE)
for i, (yr, tag, desc, c) in enumerate(phases):
    x = x0 + i * (cw + gap)
    # node
    node = rect(s, x + cw/2 - Inches(0.12), y0 - Inches(0.55), Inches(0.24), Inches(0.24), fill=c)
    rect(s, x, y0, cw, Inches(3.4), fill=PANEL, line=LINE, line_w=1)
    txt(s, x, y0 + Inches(0.25), cw, Inches(0.5),
        [[(yr, 18, WHITE, True, False)]], align=PP_ALIGN.CENTER)
    rect(s, x + cw/2 - Inches(0.95), y0 + Inches(0.85), Inches(1.9), Inches(0.5), fill=c)
    txt(s, x + cw/2 - Inches(0.95), y0 + Inches(0.88), Inches(1.9), Inches(0.45),
        [[(tag, 14, INK, True, False)]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    txt(s, x + Inches(0.22), y0 + Inches(1.55), cw - Inches(0.44), Inches(1.7),
        [[(desc, 13, MUTED, False, False)]], align=PP_ALIGN.CENTER)
txt(s, Inches(0.7), Inches(6.25), Inches(12), Inches(0.7),
    [[("Bottom line: expect ", 14, WHITE, False, False),
      ("~3 years as a no-compromise card and 5–7 years of genuinely usable high-end gaming",
       14, ACCENT, True, False),
      (" from a 2025 purchase.", 14, WHITE, False, False)]])

# ===========================================================================
# Slide 5 — What extends vs shortens its life
# ===========================================================================
s = prs.slides.add_slide(BLANK); bg(s)
kicker(s, "Risk & resilience"); accent_bar(s)
title(s, "What extends — and what shortens — its life")

rect(s, Inches(0.7), Inches(1.8), Inches(5.95), Inches(4.6), fill=PANEL, line=LINE, line_w=1)
rect(s, Inches(0.7), Inches(1.8), Inches(5.95), Inches(0.6), fill=ACCENT)
txt(s, Inches(0.7), Inches(1.83), Inches(5.95), Inches(0.55),
    [[("EXTENDS RELEVANCE", 15, INK, True, False)]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
bullets(s, Inches(1.0), Inches(2.6), Inches(5.4), Inches(3.6), [
    "DLSS 4 & future DLSS revisions adding free performance",
    "Strong raw horsepower with real 4K headroom",
    "Gaming at 1440p instead of 4K (big lifespan boost)",
    "Modern I/O — no platform bottleneck",
    "Premium Astral cooling = stable clocks & longevity",
], size=14, gap=11, dot=ACCENT)

rect(s, Inches(6.9), Inches(1.8), Inches(5.75), Inches(4.6), fill=PANEL, line=LINE, line_w=1)
rect(s, Inches(6.9), Inches(1.8), Inches(5.75), Inches(0.6), fill=WARN)
txt(s, Inches(6.9), Inches(1.83), Inches(5.75), Inches(0.55),
    [[("SHORTENS RELEVANCE", 15, INK, True, False)]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
bullets(s, Inches(7.2), Inches(2.6), Inches(5.2), Inches(3.6), [
    "16 GB VRAM ceiling as games grow more memory-hungry",
    "Heavy path-traced / ray-traced titles aging it faster",
    "Chasing 4K at max settings with no upscaling",
    "A future console generation raising the baseline",
    "Rapid AI/feature gating to newer architectures",
], size=14, gap=11, dot=WARN)

# ===========================================================================
# Slide 6 — Verdict / recommendation
# ===========================================================================
s = prs.slides.add_slide(BLANK); bg(s)
kicker(s, "Verdict", ACCENT); accent_bar(s)
title(s, "Should you count on it lasting?")

rect(s, Inches(0.7), Inches(1.8), Inches(11.95), Inches(1.5), fill=PANEL, line=ACCENT, line_w=1.5)
txt(s, Inches(1.0), Inches(1.95), Inches(11.4), Inches(1.2),
    [[("Yes — it's a safe long-term buy. ", 20, ACCENT, True, False),
      ("As a near-top-tier Blackwell card with DLSS 4, the ROG Astral RTX 5080 should "
       "stay a high-end gaming GPU for roughly 5–7 years, with its first ~3 years "
       "as a true no-compromise 4K card.", 16, WHITE, False, False)]],
    anchor=MSO_ANCHOR.MIDDLE)

txt(s, Inches(0.7), Inches(3.6), Inches(12), Inches(0.5),
    [[("Best matched to:", 16, WHITE, True, False)]])
bullets(s, Inches(0.9), Inches(4.1), Inches(11.6), Inches(2.0), [
    "4K and high-refresh 1440p gamers who want years before re-thinking an upgrade",
    "Buyers who value DLSS / frame-gen and will keep using it as titles get heavier",
    "Creators and enthusiasts who want flagship-class cooling and acoustics",
], size=15, gap=12, dot=ACCENT2)

txt(s, Inches(0.7), Inches(6.5), Inches(12), Inches(0.6),
    [[("Watch item: ", 13, WARN, True, False),
      ("the 16 GB VRAM is the single factor most likely to define its eventual "
       "ceiling — the rest of the card will outlive it.", 13, MUTED, False, True)]])

# ===========================================================================
# Slide 7 — Closing / notes
# ===========================================================================
s = prs.slides.add_slide(BLANK); bg(s)
rect(s, 0, 0, SW, Inches(0.18), fill=ACCENT)
txt(s, Inches(0.9), Inches(2.6), Inches(11.5), Inches(1.2),
    [[("Relevance is a spectrum, not a cliff.", 30, WHITE, True, False)]])
txt(s, Inches(0.9), Inches(3.7), Inches(11.5), Inches(1.4),
    [[("The RTX 5080 won't “stop working” — it gradually steps down from "
       "no-compromise 4K to a strong 1440p card. Pairing it with DLSS and a sensible "
       "resolution target is what stretches its useful life toward the longer end of "
       "the estimate.", 16, MUTED, False, False)]])
txt(s, Inches(0.9), Inches(6.4), Inches(11.5), Inches(0.5),
    [[("Estimates are projections based on historical GPU lifecycles and current "
       "trends, not guarantees.", 12, MUTED, False, True)]])

prs.save("/home/user/Test/RTX5080-Astral-Longevity.pptx")
print("Saved RTX5080-Astral-Longevity.pptx with", len(prs.slides._sldIdLst), "slides")
